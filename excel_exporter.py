import io
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from sp_api_client import TARGET_COLUMNS

# Column specific number formats matching Target specifications
COLUMN_NUMBER_FORMATS = {
    "Current Est Sales": "#,##0",
    "Estimated Monthly Revenue ($)": "$#,##0.00",
    "Profit ($)": "$#,##0.00",
    "Margin (%)": "0.00%",
    "ROI (%)": "0.00%",
    "Sales Rank": "#,##0",
    "Sub Category Sales Rank": "#,##0",
    "Cost": "$#,##0.00",
    "Break Even Sell Price": "$#,##0.00",
    "Sell Price": "$#,##0.00",
    "Buybox Landed": "$#,##0.00",
    "Lowest FBA": "$#,##0.00",
    "Lowest FBM": "$#,##0.00",
    "Referral Fee %": "0.0%",
    "Referral Fee": "$#,##0.00",
    "FBA Pick & Pack Fee ($)": "$#,##0.00",
    "Inbound Placement Fee": "$#,##0.00",
    "Fulfillment Subtotal": "$#,##0.00",
    "Cost Subtotal": "$#,##0.00",
    "Seller Proceeds": "$#,##0.00",
    "VAT %": "0.0%",
    "VAT $": "$#,##0.00",
    "Inbound Shipping Estimate": "$#,##0.00",
    "Package Weight (lb)": "0.000",
    "Total Offers": "#,##0",
    "New FBA Offers": "#,##0",
    "New FBM Offers": "#,##0",
    "BuyBox Contention": "#,##0",
    "Package Quantity": "#,##0",
    "Last Run": r"mm/dd/yyyy\ hh:mm",
    "Average Rating": "0.0",
    "Total Review Count": "#,##0",
}

INTEGER_COLUMNS = {
    "Sales Rank",
    "Sub Category Sales Rank",
    "Current Est Sales",
    "Total Offers",
    "New FBA Offers",
    "New FBM Offers",
    "BuyBox Contention",
    "Package Quantity",
    "Total Review Count",
}

# Column widths: Compact, balanced, and capped so no column is excessively wide
COLUMN_WIDTHS = {
    # Key text columns (max width 25.0)
    "Title": 25.0,
    "About this item": 24.0,
    "Customer Reviews": 20.0,
    "Wholesaller Company Name": 18.0,
    "Category": 16.0,
    "Sub Category": 20.0,
    "Brand": 14.0,
    "Manufacturer": 14.0,
    "Custom Notes": 15.0,

    # Financial & analysis columns
    "Estimated Monthly Revenue ($)": 16.0,
    "FBA Pick & Pack Fee ($)": 14.0,
    "Inbound Placement Fee": 14.0,
    "Inbound Shipping Estimate": 15.0,
    "Fulfillment Subtotal": 14.0,
    "Cost Subtotal": 13.0,
    "Seller Proceeds": 13.0,
    "Break Even Sell Price": 13.0,
    "Buybox Landed": 12.0,
    "Package Dimensions (in)": 15.0,
    "Package Weight (lb)": 12.0,
    "Current Est Sales": 12.0,
    "BuyBox Contention": 12.0,
    "Total Review Count": 12.0,
    "Referral Fee %": 12.0,
    "Referral Fee": 12.0,
    "Profit ($)": 11.5,
    "Margin (%)": 11.0,
    "ROI (%)": 11.0,
    "Sales Rank": 11.0,
    "Sub Category Sales Rank": 15.0,
    "Cost": 11.0,
    "Sell Price": 11.0,
    "Lowest FBA": 11.0,
    "Lowest FBM": 11.0,
    "Average Rating": 11.0,
    "VAT %": 10.0,
    "VAT $": 10.5,
    "Total Offers": 10.5,
    "New FBA Offers": 10.5,
    "New FBM Offers": 10.5,
    "Package Quantity": 11.0,

    # Identifiers & metadata
    "Product ID": 13.0,
    "ASIN": 11.5,
    "Parent ASIN": 11.5,
    "Model": 12.0,
    "Part Number": 12.0,
    "Color": 11.0,
    "Is Hazmat": 10.5,
    "Brand Restriction": 12.0,
    "Selling Restriction": 14.0,
    "Selling Restriction Link": 15.0,
    "Is Top Level Category": 13.0,
    "Last Run": 15.0,
    "Favorite": 9.0,
    "Amazon Link": 13.0,
    "Keepa": 10.0,
    "CamelCC": 10.0,

    # Product image links (compact)
    "Product Image link 1": 14.0,
    "Product Image link 2": 14.0,
    "Product Image link 3": 14.0,
    "Product Image link 4": 14.0,
    "Product Image link 5": 14.0,
}
DEFAULT_COLUMN_WIDTH = 12.0

# Right-aligned columns: financial, currency, metrics, counts, and percentages
RIGHT_ALIGNED_COLUMNS = {
    "Current Est Sales",
    "Estimated Monthly Revenue ($)",
    "Profit ($)",
    "Margin (%)",
    "ROI (%)",
    "Sales Rank",
    "Sub Category Sales Rank",
    "Cost",
    "Break Even Sell Price",
    "Sell Price",
    "Buybox Landed",
    "Lowest FBA",
    "Lowest FBM",
    "Referral Fee %",
    "Referral Fee",
    "FBA Pick & Pack Fee ($)",
    "Inbound Placement Fee",
    "Fulfillment Subtotal",
    "Cost Subtotal",
    "Seller Proceeds",
    "VAT %",
    "VAT $",
    "Inbound Shipping Estimate",
    "Package Weight (lb)",
    "Total Offers",
    "New FBA Offers",
    "New FBM Offers",
    "BuyBox Contention",
    "Package Quantity",
    "Average Rating",
    "Total Review Count",
}

# Center-aligned columns: codes, IDs, booleans/flags, dates
CENTER_ALIGNED_COLUMNS = {
    "ASIN",
    "Parent ASIN",
    "Favorite",
    "Product ID",
    "Brand Restriction",
    "Selling Restriction",
    "Is Top Level Category",
    "Package Dimensions (in)",
    "Color",
    "Model",
    "Part Number",
    "Is Hazmat",
    "Keepa",
    "CamelCC",
    "Last Run",
}


def export_to_target_excel(rows_or_df, output_path=None):
    """
    Exports a list of row dictionaries or a pandas DataFrame into a beautifully formatted,
    executive-grade Excel (.xlsx) workbook.

    Key design rules:
    - Fixed row heights: Header is 24pt, every data row is strictly 20pt (single-line height).
    - Compact column widths: Balanced and capped (max width 25.0) to avoid overly wide sheets.
    - No multi-line expansion: wrap_text=False across all cells; linebreaks in descriptions
      are cleanly converted to single-line separator (' | ') to preserve text without swelling rows.
    - Professional visual theme: Corporate navy header, subtle zebra striping, clean light borders,
      gridlines enabled, and auto-filters with top-row freeze panes.

    Returns:
        bytes: If output_path is None, returns Excel file binary data (for Streamlit downloads).
        output_path: If output_path is provided, writes to the file path and returns output_path.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog Analysis"

    # Display grid lines and freeze header row
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    # Styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    header_border = Border(
        left=Side(style="thin", color="17365D"),
        right=Side(style="thin", color="17365D"),
        top=Side(style="thin", color="17365D"),
        bottom=Side(style="medium", color="17365D"),
    )

    cell_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    data_font = Font(name="Calibri", size=10, bold=False, color="1E293B")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # 1. Write Header Row (fixed height 24pt)
    ws.row_dimensions[1].height = 24.0
    for col_idx, col_name in enumerate(TARGET_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = header_border

        # Set column width
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(col_name, DEFAULT_COLUMN_WIDTH)

    # Convert DataFrame to records if passed
    if hasattr(rows_or_df, "to_dict"):
        records = rows_or_df.to_dict(orient="records")
    else:
        records = list(rows_or_df)

    # 2. Write Data Rows (each strictly 20pt height = height of 1 line)
    for row_idx, record in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20.0
        is_zebra = (row_idx % 2 == 1)

        for col_idx, col_name in enumerate(TARGET_COLUMNS, 1):
            val = record.get(col_name, None)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border

            if is_zebra:
                cell.fill = zebra_fill

            # Determine horizontal alignment
            if col_name in RIGHT_ALIGNED_COLUMNS:
                h_align = "right"
            elif col_name in CENTER_ALIGNED_COLUMNS:
                h_align = "center"
            else:
                h_align = "left"

            # Check if value is None or pandas NaN
            if val is None or (isinstance(val, float) and pd.isna(val)) or str(val).strip().lower() == "nan":
                cell.value = None
            else:
                # String conversion and line-break flattening for single-line consistency
                if col_name in (
                    "Product ID", "ASIN", "Parent ASIN", "Wholesaller Company Name",
                    "Customer Reviews", "About this item", "Title", "Brand", "Category", "Sub Category"
                ) and not isinstance(val, str):
                    val = str(val)

                if isinstance(val, str):
                    # Replace internal linebreaks so Excel never expands row height
                    if "\n" in val or "\r" in val:
                        cleaned_lines = [
                            p.strip() for p in val.replace("\r\n", "\n").replace("\r", "\n").split("\n")
                            if p.strip()
                        ]
                        val = " | ".join(cleaned_lines)

                # Pure integer columns (Sales Rank, Offers, Reviews, etc.)
                if col_name in INTEGER_COLUMNS:
                    try:
                        val = int(round(float(val)))
                    except (ValueError, TypeError):
                        pass
                # Float columns - ensure max 3 decimal fraction points
                elif isinstance(val, float):
                    val = round(val, 3)

                cell.value = val

                # Apply number formatting if specified
                if col_name in COLUMN_NUMBER_FORMATS:
                    cell.number_format = COLUMN_NUMBER_FORMATS[col_name]

            # Enforce single line display without text wrapping
            cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=False)

    # 3. Add Autofilter
    total_rows = len(records) + 1
    last_col_letter = get_column_letter(len(TARGET_COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col_letter}{total_rows}"

    if output_path:
        wb.save(output_path)
        return output_path
    else:
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

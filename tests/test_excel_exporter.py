"""
Tests for Module 2: Target 52-Column Excel Exporter.
"""

import io
import unittest
import pandas as pd
import openpyxl
from modules.analyzer.excel_exporter import export_to_target_excel
from modules.analyzer.sp_api_client import TARGET_COLUMNS


class TestExcelExporter(unittest.TestCase):
    def test_export_to_target_excel(self):
        # Create minimal DataFrame with TARGET_COLUMNS
        row = {col: "Sample" for col in TARGET_COLUMNS}
        row["Cost"] = 15.00
        row["Sell Price"] = 29.99
        row["Profit ($)"] = 5.50
        row["ROI (%)"] = 0.3667
        row["Margin (%)"] = 0.1834
        row["Sales Rank"] = 12000

        df = pd.DataFrame([row])
        excel_bytes = export_to_target_excel(df)

        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 0)

        # Parse with openpyxl to verify worksheet structure
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        self.assertIsNotNone(ws)
        # Header count must match
        header_vals = [cell.value for cell in ws[1]]
        self.assertEqual(len(header_vals), len(TARGET_COLUMNS))
        self.assertEqual(header_vals[0], TARGET_COLUMNS[0])

    def test_quick_search_export_to_excel(self):
        from modules.analyzer.sp_api_client import fetch_amazon_item_data, build_full_target_row
        item = fetch_amazon_item_data(
            upc="794437466810",
            asin="B01N1S89M0",
            title="SheaMoisture Daily Hydration Conditioner 13 oz",
            cost=5.25,
            credentials={},
            marketplace="US",
            is_demo=True,
        )
        row = build_full_target_row(item_data=item, cost=5.25, wholesaler_company_name="Acme Wholesale Supplies LLC")
        df = pd.DataFrame([row])
        excel_bytes = export_to_target_excel(df)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # Header row + 1 data row
        self.assertEqual(ws.max_column, len(TARGET_COLUMNS))  # Exact target schema columns

    def test_analyzing_software_name_filled_with_internal_amz(self):
        """Verify that Analyzing Software Name is automatically filled with 'InternalAMZ' before generating Excel."""
        # 1. Test when the column is completely omitted in input DataFrame
        df_without_col = pd.DataFrame([{"ASIN": "B01N1S89M0", "Cost": 10.0, "Sell Price": 25.0}])
        excel_bytes = export_to_target_excel(df_without_col)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        self.assertIn("Analyzing Software Name", headers)
        col_idx = headers.index("Analyzing Software Name") + 1

        val = ws.cell(row=2, column=col_idx).value
        self.assertEqual(val, "InternalAMZ")

        # 2. Test when the column exists but has None / empty values
        df_with_empty = pd.DataFrame([{"ASIN": "B01N1S89M0", "Analyzing Software Name": None}])
        excel_bytes2 = export_to_target_excel(df_with_empty)
        wb2 = openpyxl.load_workbook(io.BytesIO(excel_bytes2))
        ws2 = wb2.active

        val2 = ws2.cell(row=2, column=col_idx).value
        self.assertEqual(val2, "InternalAMZ")

        # 3. Test when input is a list of dicts
        dict_records = [{"ASIN": "B01N1S89M0"}]
        excel_bytes3 = export_to_target_excel(dict_records)
        wb3 = openpyxl.load_workbook(io.BytesIO(excel_bytes3))
        ws3 = wb3.active
        val3 = ws3.cell(row=2, column=col_idx).value
        self.assertEqual(val3, "InternalAMZ")



if __name__ == "__main__":
    unittest.main()

